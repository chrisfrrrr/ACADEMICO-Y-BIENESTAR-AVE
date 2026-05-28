import io
import re
import zipfile
from datetime import datetime, date
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import plotly.express as px
import requests
import streamlit as st

try:
    import gspread
    from google.oauth2.service_account import Credentials
except Exception:
    gspread = None
    Credentials = None

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
except Exception:
    Document = None

APP_NAME = "Sistema PRO de Seguimiento y Derivación Académica AVE"
RIESGO_ORDEN = {"Bajo": 1, "Moderado": 2, "Alto": 3}
RIESGO_COLOR = {"Bajo": "#0F766E", "Moderado": "#B7791F", "Alto": "#B91C1C"}

SHEETS = {
    "Estudiantes": ["carne", "nombre", "correo", "telefono", "carrera", "trimestre", "curso", "seccion", "canvas_user_id", "asesor_bienestar"],
    "Asesores_Bienestar": ["id_asesor", "nombre", "correo", "telefono", "observaciones"],
    "Historial_Estudiantes": ["fecha", "id_consulta", "carne", "nombre", "correo", "curso", "seccion", "semana", "actividades_pct", "promedio", "entregas_tarde", "semanas_sin_entregas", "ingresos_semana", "dias_inactivo", "horas_respuesta", "riesgo", "riesgo_anterior", "cambio", "motivo_detectado", "asesor_academico"],
    "Derivaciones": ["id_derivacion", "fecha", "carne", "nombre", "correo", "curso", "seccion", "riesgo", "prioridad", "asesor_bienestar", "correo_bienestar", "motivo", "acciones_previas", "observaciones", "estado_derivacion", "asesor_academico"],
    "Mensajes_Enviados": ["fecha", "carne", "nombre", "correo", "curso", "riesgo", "tipo_mensaje", "mensaje_generado", "enviado_canvas", "asesor_academico"],
    "Consultas_Canvas": ["id_consulta", "fecha_consulta", "asesor_academico", "curso", "semana", "total_estudiantes", "bajo", "moderado", "alto", "fuente_datos"],
    "Configuracion": ["parametro", "valor"],
}

# -----------------------------------------------------------------------------
# UI
# -----------------------------------------------------------------------------
st.set_page_config(page_title="Seguimiento AVE PRO", page_icon="🎓", layout="wide")
st.markdown("""
<style>
.main .block-container{padding-top:1.1rem; padding-bottom:2rem;}
.metric-card{border-radius:18px;padding:18px;background:#ffffff;border:1px solid #e5e7eb;box-shadow:0 8px 22px rgba(15,23,42,.06)}
.small-muted{font-size:0.87rem;color:#64748b}.risk-low{color:#0F766E;font-weight:700}.risk-mid{color:#B7791F;font-weight:700}.risk-high{color:#B91C1C;font-weight:700}
.section-title{font-size:1.2rem;font-weight:800;margin-top:1rem;color:#0F172A}.pill{border-radius:999px;padding:4px 10px;background:#f1f5f9;color:#334155;font-size:.85rem}
</style>
""", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# Helpers DB
# -----------------------------------------------------------------------------
def empty_db() -> Dict[str, pd.DataFrame]:
    return {name: pd.DataFrame(columns=cols) for name, cols in SHEETS.items()}


def normalize_cols(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip().lower().replace(" ", "_").replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u") for c in df.columns]
    return df


def load_excel_db(uploaded_file) -> Dict[str, pd.DataFrame]:
    db = empty_db()
    if uploaded_file is None:
        return db
    xls = pd.ExcelFile(uploaded_file)
    for sheet in xls.sheet_names:
        key = next((s for s in SHEETS if s.lower() == sheet.lower()), sheet)
        if key in SHEETS:
            df = pd.read_excel(xls, sheet_name=sheet)
            df = normalize_cols(df)
            for col in SHEETS[key]:
                if col not in df.columns:
                    df[col] = np.nan
            db[key] = df[SHEETS[key]]
    return db


def export_db_excel(db: Dict[str, pd.DataFrame]) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        for name, cols in SHEETS.items():
            df = db.get(name, pd.DataFrame(columns=cols)).copy()
            for c in cols:
                if c not in df.columns:
                    df[c] = np.nan
            df = df[cols]
            df.to_excel(writer, index=False, sheet_name=name)
            workbook = writer.book
            ws = writer.sheets[name]
            header_fmt = workbook.add_format({"bold": True, "bg_color": "#0B1F3A", "font_color": "#FFFFFF", "border": 1})
            body_fmt = workbook.add_format({"text_wrap": True, "valign": "top"})
            for i, col in enumerate(df.columns):
                ws.write(0, i, col, header_fmt)
                width = min(max(12, int(df[col].astype(str).str.len().quantile(.75) if len(df) else len(col)) + 3), 36)
                ws.set_column(i, i, width, body_fmt)
            ws.freeze_panes(1, 0)
    return output.getvalue()


def append_rows(db: Dict[str, pd.DataFrame], sheet: str, rows: List[Dict]) -> Dict[str, pd.DataFrame]:
    if not rows:
        return db
    new_df = pd.DataFrame(rows)
    for col in SHEETS[sheet]:
        if col not in new_df.columns:
            new_df[col] = np.nan
    db[sheet] = pd.concat([db.get(sheet, pd.DataFrame(columns=SHEETS[sheet])), new_df[SHEETS[sheet]]], ignore_index=True)
    return db

# -----------------------------------------------------------------------------
# Canvas Client
# -----------------------------------------------------------------------------
class CanvasClient:
    def __init__(self, base_url: str, token: str):
        self.base_url = base_url.rstrip("/")
        self.headers = {"Authorization": f"Bearer {token}"}

    def _get_paginated(self, path: str, params: Optional[dict] = None) -> List[dict]:
        url = f"{self.base_url}/api/v1{path}"
        params = params or {}
        params.setdefault("per_page", 100)
        out = []
        while url:
            r = requests.get(url, headers=self.headers, params=params, timeout=35)
            r.raise_for_status()
            data = r.json()
            out.extend(data if isinstance(data, list) else [data])
            url = None
            if "next" in r.links:
                url = r.links["next"]["url"]
                params = None
        return out

    def validate(self) -> Tuple[bool, str]:
        try:
            me = requests.get(f"{self.base_url}/api/v1/users/self", headers=self.headers, timeout=20)
            if me.ok:
                data = me.json()
                return True, data.get("name", "Usuario validado")
            return False, f"Canvas respondió {me.status_code}: {me.text[:180]}"
        except Exception as e:
            return False, str(e)

    def courses(self) -> pd.DataFrame:
        data = self._get_paginated("/courses", {"enrollment_state": "active"})
        return pd.DataFrame([{"id": c.get("id"), "name": c.get("name"), "course_code": c.get("course_code")} for c in data])

    def users(self, course_id: str) -> pd.DataFrame:
        data = self._get_paginated(f"/courses/{course_id}/users", {"enrollment_type[]": "student"})
        return pd.DataFrame([{"canvas_user_id": u.get("id"), "nombre": u.get("name"), "correo": u.get("email"), "login_id": u.get("login_id")} for u in data])

    def assignments(self, course_id: str) -> pd.DataFrame:
        data = self._get_paginated(f"/courses/{course_id}/assignments", {"include[]": ["submission"]})
        rows = []
        for a in data:
            rows.append({"assignment_id": a.get("id"), "nombre_actividad": a.get("name"), "puntos": a.get("points_possible") or 0, "due_at": a.get("due_at"), "published": a.get("published")})
        return pd.DataFrame(rows)

    def submissions(self, course_id: str, assignment_id: str) -> pd.DataFrame:
        data = self._get_paginated(f"/courses/{course_id}/assignments/{assignment_id}/submissions", {"include[]": ["user"]})
        rows = []
        for s in data:
            u = s.get("user") or {}
            rows.append({"assignment_id": assignment_id, "canvas_user_id": s.get("user_id"), "nombre": u.get("name"), "submitted_at": s.get("submitted_at"), "late": s.get("late"), "missing": s.get("missing"), "score": s.get("score"), "workflow_state": s.get("workflow_state")})
        return pd.DataFrame(rows)

    def send_message(self, recipients: List[str], subject: str, body: str) -> Tuple[bool, str]:
        payload = {"recipients[]": recipients, "subject": subject, "body": body, "force_new": True}
        try:
            r = requests.post(f"{self.base_url}/api/v1/conversations", headers=self.headers, data=payload, timeout=25)
            if r.ok:
                return True, "Mensaje enviado por Canvas."
            return False, f"No se pudo enviar. Canvas respondió {r.status_code}: {r.text[:180]}"
        except Exception as e:
            return False, str(e)

# -----------------------------------------------------------------------------
# Risk logic and messages
# -----------------------------------------------------------------------------
def classify_student(row: pd.Series) -> Tuple[str, str]:
    pct = pd.to_numeric(row.get("actividades_pct"), errors="coerce")
    prom = pd.to_numeric(row.get("promedio"), errors="coerce")
    tarde = pd.to_numeric(row.get("entregas_tarde"), errors="coerce")
    sin_ent = pd.to_numeric(row.get("semanas_sin_entregas"), errors="coerce")
    ingresos = pd.to_numeric(row.get("ingresos_semana"), errors="coerce")
    inactivo = pd.to_numeric(row.get("dias_inactivo"), errors="coerce")
    horas_resp = pd.to_numeric(row.get("horas_respuesta"), errors="coerce")
    motivos = []

    high_flags = 0
    mod_flags = 0
    if pd.notna(pct):
        if pct < 50: high_flags += 1; motivos.append("bajo cumplimiento de actividades semanales")
        elif pct < 80: mod_flags += 1; motivos.append("cumplimiento parcial de actividades")
    if pd.notna(prom):
        if prom < 59: high_flags += 1; motivos.append("promedio inferior al mínimo esperado")
        elif prom < 70: mod_flags += 1; motivos.append("promedio académico en rango de alerta")
    if pd.notna(tarde):
        if tarde >= 4: high_flags += 1; motivos.append("entregas pendientes o tardías recurrentes")
        elif tarde >= 2: mod_flags += 1; motivos.append("entregas irregulares o tardías")
    if pd.notna(sin_ent) and sin_ent >= 2:
        high_flags += 1; motivos.append("sin entregas por dos o más semanas")
    if pd.notna(ingresos):
        if ingresos == 0: high_flags += 1; motivos.append("sin ingresos semanales a Canvas")
        elif ingresos <= 2: mod_flags += 1; motivos.append("baja frecuencia de ingreso a Canvas")
    if pd.notna(inactivo):
        if inactivo >= 6: high_flags += 1; motivos.append("inactividad total en Canvas por seis o más días")
        elif inactivo >= 3: mod_flags += 1; motivos.append("actividad limitada en Canvas")
    if pd.notna(horas_resp):
        if horas_resp >= 120: high_flags += 1; motivos.append("no responde comunicaciones por cinco o más días")
        elif horas_resp >= 48: mod_flags += 1; motivos.append("responde comunicaciones con retraso")

    if high_flags >= 1 and (high_flags + mod_flags) >= 2:
        return "Alto", ", ".join(dict.fromkeys(motivos)) or "riesgo académico alto"
    if high_flags >= 2:
        return "Alto", ", ".join(dict.fromkeys(motivos)) or "riesgo académico alto"
    if mod_flags >= 1 or high_flags == 1:
        return "Moderado", ", ".join(dict.fromkeys(motivos)) or "señales académicas de alerta"
    return "Bajo", "desempeño académico estable"


def compare_risk(prev: Optional[str], current: str) -> str:
    if not prev or pd.isna(prev): return "Sin registro previo"
    if RIESGO_ORDEN.get(current, 0) < RIESGO_ORDEN.get(prev, 0): return "Mejora"
    if RIESGO_ORDEN.get(current, 0) > RIESGO_ORDEN.get(prev, 0): return "Empeora"
    return "Sin cambio"


def generate_message(row: pd.Series, asesor: str, horario: str, canal: str) -> str:
    nombre = row.get("nombre", "estudiante")
    riesgo = row.get("riesgo", "Bajo")
    motivo = row.get("motivo_detectado", "tu avance académico")
    if riesgo == "Bajo":
        return f"""Hola, {nombre}:\n\nEspero que te encontrés muy bien. Al revisar tu avance de esta semana, observo que mantenés un seguimiento adecuado del curso, especialmente en el cumplimiento de actividades, participación y desempeño general.\n\nTe felicito por la constancia que has demostrado y te animo a continuar con ese ritmo de trabajo. Recordá que cualquier duda o situación que necesités conversar podés escribirme por este medio.\n\nSaludos cordiales,\n{asesor}"""
    if riesgo == "Moderado":
        return f"""Hola, {nombre}:\n\nEspero que te encontrés bien. Al revisar tu avance académico de esta semana, identifiqué algunos aspectos que requieren atención, principalmente relacionados con {motivo}.\n\nEl propósito de este mensaje es brindarte acompañamiento oportuno y apoyarte para que puedas retomar el ritmo del curso antes de que la situación afecte de forma significativa tu desempeño académico. Te recomiendo revisar las actividades pendientes y organizar tus tiempos de entrega.\n\nQuedo atento a tu respuesta para conocer si existe alguna situación en la que podamos orientarte o apoyarte.\n\nSaludos cordiales,\n{asesor}"""
    return f"""Hola, {nombre}:\n\nEspero que te encontrés bien. Te escribo porque, al revisar tu participación y avance en Canvas, se identificaron señales importantes de riesgo académico relacionadas con {motivo}.\n\nEsta situación requiere atención prioritaria, ya que puede afectar tu continuidad y desempeño en el curso. Por ello, me gustaría brindarte un espacio de seguimiento más cercano.\n\nTe propongo que podamos reunirnos o comunicarnos en el siguiente horario: {horario}. Canal de atención: {canal}.\n\nQuedo atento a tu confirmación para poder apoyarte de manera oportuna.\n\nSaludos cordiales,\n{asesor}"""

# -----------------------------------------------------------------------------
# Derivation documents
# -----------------------------------------------------------------------------
def add_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text) if text is not None else "")
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_derivation_doc(row: pd.Series, asesor_academico: str, asesor_bienestar: str, observaciones: str, acciones: str) -> bytes:
    if Document is None:
        raise RuntimeError("python-docx no está instalado.")
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55); sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = title.add_run("AVE UVG\nFormato de Derivación Académica")
    rr.bold = True; rr.font.size = Pt(16); rr.font.color.rgb = RGBColor(11, 31, 58)
    risk = row.get("riesgo", "")
    if risk == "Alto":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("PRIORIDAD ALTA - RIESGO DE ABANDONO")
        run.bold = True; run.font.size = Pt(12); run.font.color.rgb = RGBColor(185, 28, 28)
    doc.add_paragraph(f"Fecha de derivación: {date.today().strftime('%d/%m/%Y')}")

    sections = [
        ("1. Datos del estudiante", [("Nombre del estudiante", row.get("nombre", "")), ("Carné", row.get("carne", "")), ("Carrera", row.get("carrera", "")), ("Trimestre", row.get("trimestre", "")), ("Correo electrónico", row.get("correo", "")), ("Teléfono", row.get("telefono", ""))]),
        ("2. Datos del remitente", [("Nombre del asesor académico", asesor_academico), ("Asesor de bienestar receptor", asesor_bienestar), ("Curso", row.get("curso", "")), ("Sección", row.get("seccion", ""))]),
        ("3. Motivo de la derivación", [("Nivel de riesgo", risk), ("Motivo detectado", row.get("motivo_detectado", ""))]),
        ("4. Descripción breve del caso", [("Descripción", f"El estudiante presenta nivel de riesgo {risk.lower()} debido a {row.get('motivo_detectado', 'indicadores académicos de alerta')}. Se recomienda seguimiento por parte del área de bienestar estudiantil conforme al protocolo institucional.")]),
        ("5. Acciones previas realizadas", [("Acciones", acciones)]),
        ("6. Observaciones adicionales", [("Observaciones", observaciones)]),
        ("7. Nivel de prioridad", [("Prioridad", "Alta (riesgo de abandono)" if risk == "Alto" else "Media (requiere apoyo pronto)")]),
    ]
    for head, items in sections:
        hp = doc.add_paragraph()
        run = hp.add_run(head)
        run.bold = True; run.font.color.rgb = RGBColor(11, 31, 58)
        table = doc.add_table(rows=len(items), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for i, (k, v) in enumerate(items):
            add_cell_text(table.cell(i,0), k, bold=True, color="#0B1F3A")
            add_cell_text(table.cell(i,1), v)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# -----------------------------------------------------------------------------
# Data processing
# -----------------------------------------------------------------------------
def standardize_analysis_input(df: pd.DataFrame, db: Dict[str, pd.DataFrame], curso: str, semana: int, asesor: str) -> pd.DataFrame:
    df = normalize_cols(df)
    aliases = {
        "nombre_del_estudiante": "nombre", "estudiante": "nombre", "sis_user_id": "carne", "id": "carne",
        "mail": "correo", "email": "correo", "score": "promedio", "calificacion": "promedio",
        "actividades_completadas_%": "actividades_pct", "porcentaje_actividades": "actividades_pct", "avance": "actividades_pct",
        "late": "entregas_tarde", "tardias": "entregas_tarde", "missing": "semanas_sin_entregas",
        "ingresos": "ingresos_semana", "dias_sin_actividad": "dias_inactivo", "respuesta_horas": "horas_respuesta"
    }
    df = df.rename(columns={c: aliases.get(c, c) for c in df.columns})
    for col in ["carne", "nombre", "correo", "telefono", "carrera", "trimestre", "curso", "seccion", "canvas_user_id", "asesor_bienestar", "actividades_pct", "promedio", "entregas_tarde", "semanas_sin_entregas", "ingresos_semana", "dias_inactivo", "horas_respuesta"]:
        if col not in df.columns:
            df[col] = np.nan
    df["curso"] = df["curso"].fillna(curso)
    df.loc[df["curso"].astype(str).str.strip().eq(""), "curso"] = curso
    est = db.get("Estudiantes", pd.DataFrame())
    if not est.empty and "carne" in df.columns:
        base_cols = [c for c in ["carne", "telefono", "carrera", "trimestre", "seccion", "asesor_bienestar", "canvas_user_id"] if c in est.columns]
        df = df.merge(est[base_cols].drop_duplicates("carne"), on="carne", how="left", suffixes=("", "_base"))
        for c in ["telefono", "carrera", "trimestre", "seccion", "asesor_bienestar", "canvas_user_id"]:
            if f"{c}_base" in df.columns:
                df[c] = df[c].combine_first(df[f"{c}_base"])
                df.drop(columns=[f"{c}_base"], inplace=True)
    risks, motives = [], []
    for _, row in df.iterrows():
        r, m = classify_student(row)
        risks.append(r); motives.append(m)
    df["riesgo"] = risks
    df["motivo_detectado"] = motives
    hist = db.get("Historial_Estudiantes", pd.DataFrame())
    prev_map = {}
    if not hist.empty and "carne" in hist.columns:
        h = hist.dropna(subset=["carne"]).copy()
        h["fecha_sort"] = pd.to_datetime(h.get("fecha"), errors="coerce")
        h = h.sort_values("fecha_sort").drop_duplicates("carne", keep="last")
        prev_map = h.set_index("carne")["riesgo"].to_dict()
    df["riesgo_anterior"] = df["carne"].map(prev_map)
    df["cambio"] = [compare_risk(p, c) for p, c in zip(df["riesgo_anterior"], df["riesgo"])]
    df["semana"] = semana
    df["asesor_academico"] = asesor
    return df

# -----------------------------------------------------------------------------
# Google Sheets optional
# -----------------------------------------------------------------------------
def read_gsheet(spreadsheet_url: str, json_bytes: bytes) -> Dict[str, pd.DataFrame]:
    if gspread is None or Credentials is None:
        raise RuntimeError("gspread/google-auth no están instalados.")
    scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
    import json
    creds = Credentials.from_service_account_info(json.loads(json_bytes.decode("utf-8")), scopes=scopes)
    gc = gspread.authorize(creds)
    sh = gc.open_by_url(spreadsheet_url)
    db = empty_db()
    for name in SHEETS:
        try:
            ws = sh.worksheet(name)
            records = ws.get_all_records()
            db[name] = normalize_cols(pd.DataFrame(records)) if records else pd.DataFrame(columns=SHEETS[name])
        except Exception:
            pass
    return db

# -----------------------------------------------------------------------------
# Session state init
# -----------------------------------------------------------------------------
if "db" not in st.session_state:
    st.session_state.db = empty_db()
if "analysis_df" not in st.session_state:
    st.session_state.analysis_df = pd.DataFrame()
if "canvas_client" not in st.session_state:
    st.session_state.canvas_client = None

# -----------------------------------------------------------------------------
# Sidebar configuration
# -----------------------------------------------------------------------------
st.sidebar.image("https://dummyimage.com/480x100/0B1F3A/ffffff.png&text=AVE+Seguimiento+PRO", use_column_width=True)
st.sidebar.header("Configuración")
asesor_academico = st.sidebar.text_input("Nombre del asesor académico", value="Asesor Académico")
horario_atencion = st.sidebar.text_input("Horario de atención para riesgo alto", value="martes y jueves de 16:00 a 17:00")
canal_atencion = st.sidebar.text_input("Canal de atención", value="Bandeja de entrada de Canvas / enlace institucional")
semana_analisis = st.sidebar.selectbox("Semana de análisis", [1, 2, 3, 4, 5], index=0)
curso_manual = st.sidebar.text_input("Nombre del curso", value="Curso AVE")

st.title(APP_NAME)
st.caption("Clasificación de riesgo, historial académico, mensajes preventivos y derivaciones a bienestar con base de datos en Excel/Google Sheets.")

tabs = st.tabs(["🏠 Inicio", "🗂️ Base de datos", "🔌 Canvas / Datos", "📊 Dashboard", "👤 Estudiante", "✉️ Mensajes", "📌 Derivaciones", "⬇️ Exportar"])

# -----------------------------------------------------------------------------
# Inicio
# -----------------------------------------------------------------------------
with tabs[0]:
    st.markdown("### Propósito")
    st.write("Esta aplicación apoya el primer filtro del asesor académico mediante el análisis de actividades, calificaciones, entregas, actividad en Canvas y respuesta a comunicaciones. A partir de estos datos clasifica a cada estudiante en riesgo bajo, moderado o alto, genera acciones preventivas y registra historial para evitar duplicidad de derivaciones.")
    c1, c2, c3 = st.columns(3)
    c1.info("**Riesgo bajo**\n\nRetroalimentación cálida y seguimiento académico regular.")
    c2.warning("**Riesgo moderado**\n\nMensaje de preocupación, seguimiento y posible derivación a bienestar.")
    c3.error("**Riesgo alto**\n\nAtención prioritaria, horario de apoyo y derivación con distintivo de prioridad.")
    st.markdown("### Flujo recomendado")
    st.write("1. Cargar base de datos. 2. Conectar Canvas o cargar reporte. 3. Ejecutar análisis. 4. Revisar dashboard. 5. Generar mensajes. 6. Seleccionar derivaciones. 7. Exportar base actualizada.")

# -----------------------------------------------------------------------------
# Base de datos
# -----------------------------------------------------------------------------
with tabs[1]:
    st.markdown("### Cargar base de datos")
    uploaded_db = st.file_uploader("Cargar Excel de base de datos", type=["xlsx"], key="db_upload")
    col_a, col_b = st.columns(2)
    with col_a:
        if st.button("Usar Excel cargado", type="primary"):
            st.session_state.db = load_excel_db(uploaded_db)
            st.success("Base de datos cargada correctamente.")
    with col_b:
        if st.button("Crear base vacía"):
            st.session_state.db = empty_db()
            st.success("Base vacía creada en memoria.")

    with st.expander("Conectar Google Sheets opcional"):
        gs_url = st.text_input("URL de Google Sheets")
        gs_json = st.file_uploader("Credenciales JSON de cuenta de servicio", type=["json"], key="gs_json")
        if st.button("Leer Google Sheets"):
            try:
                st.session_state.db = read_gsheet(gs_url, gs_json.getvalue())
                st.success("Google Sheets leído correctamente.")
            except Exception as e:
                st.error(f"No se pudo leer Google Sheets: {e}")

    st.markdown("### Hojas detectadas")
    selected_sheet = st.selectbox("Vista previa", list(SHEETS.keys()))
    st.dataframe(st.session_state.db.get(selected_sheet, pd.DataFrame()), use_container_width=True, height=320)

# -----------------------------------------------------------------------------
# Canvas / Datos
# -----------------------------------------------------------------------------
with tabs[2]:
    st.markdown("### Conexión con Canvas")
    url = st.text_input("URL de Canvas", placeholder="https://uvg.instructure.com")
    token = st.text_input("Token de Canvas", type="password")
    col1, col2 = st.columns([1, 2])
    with col1:
        if st.button("Validar token"):
            try:
                client = CanvasClient(url, token)
                ok, msg = client.validate()
                if ok:
                    st.session_state.canvas_client = client
                    st.success(f"Conexión validada: {msg}")
                else:
                    st.error(msg)
            except Exception as e:
                st.error(str(e))
    with col2:
        st.info("La extracción depende de permisos del token. Si alguna métrica no está disponible, use la carga manual de reporte.")

    st.markdown("### Cargar reporte manual")
    report_file = st.file_uploader("Cargar reporte CSV o Excel con columnas de indicadores", type=["csv", "xlsx"], key="report")
    if report_file:
        if report_file.name.lower().endswith("csv"):
            raw = pd.read_csv(report_file)
        else:
            raw = pd.read_excel(report_file)
        st.write("Vista previa del reporte cargado")
        st.dataframe(raw.head(20), use_container_width=True)
        if st.button("Ejecutar análisis con reporte cargado", type="primary"):
            analyzed = standardize_analysis_input(raw, st.session_state.db, curso_manual, semana_analisis, asesor_academico)
            st.session_state.analysis_df = analyzed
            idc = datetime.now().strftime("C%Y%m%d%H%M%S")
            rows = []
            for _, r in analyzed.iterrows():
                d = r.to_dict(); d["fecha"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S"); d["id_consulta"] = idc
                rows.append(d)
            append_rows(st.session_state.db, "Historial_Estudiantes", rows)
            counts = analyzed["riesgo"].value_counts().to_dict()
            append_rows(st.session_state.db, "Consultas_Canvas", [{"id_consulta": idc, "fecha_consulta": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "asesor_academico": asesor_academico, "curso": curso_manual, "semana": semana_analisis, "total_estudiantes": len(analyzed), "bajo": counts.get("Bajo", 0), "moderado": counts.get("Moderado", 0), "alto": counts.get("Alto", 0), "fuente_datos": "reporte manual"}])
            st.success("Análisis ejecutado y registrado en historial.")

    with st.expander("Extracción básica desde Canvas"):
        course_id = st.text_input("ID del curso Canvas")
        if st.button("Obtener estudiantes del curso"):
            try:
                client = st.session_state.canvas_client or CanvasClient(url, token)
                users = client.users(course_id)
                users["curso"] = curso_manual
                users["actividades_pct"] = np.nan; users["promedio"] = np.nan; users["entregas_tarde"] = np.nan; users["semanas_sin_entregas"] = np.nan; users["ingresos_semana"] = np.nan; users["dias_inactivo"] = np.nan; users["horas_respuesta"] = np.nan
                st.session_state.analysis_df = standardize_analysis_input(users, st.session_state.db, curso_manual, semana_analisis, asesor_academico)
                st.success("Estudiantes obtenidos. Complete métricas faltantes con reporte manual si Canvas no las expone.")
                st.dataframe(st.session_state.analysis_df, use_container_width=True)
            except Exception as e:
                st.error(f"No se pudo obtener información: {e}")

# -----------------------------------------------------------------------------
# Dashboard
# -----------------------------------------------------------------------------
with tabs[3]:
    df = st.session_state.analysis_df.copy()
    if df.empty:
        st.warning("Primero ejecutá un análisis desde la pestaña Canvas / Datos.")
    else:
        f1, f2 = st.columns(2)
        risk_filter = f1.multiselect("Filtrar por riesgo", ["Bajo", "Moderado", "Alto"], default=["Bajo", "Moderado", "Alto"])
        section_filter = f2.multiselect("Filtrar por sección", sorted(df["seccion"].dropna().astype(str).unique()))
        dff = df[df["riesgo"].isin(risk_filter)]
        if section_filter:
            dff = dff[dff["seccion"].astype(str).isin(section_filter)]
        total = len(dff)
        bajo = int((dff["riesgo"] == "Bajo").sum()); mod = int((dff["riesgo"] == "Moderado").sum()); alto = int((dff["riesgo"] == "Alto").sum())
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Estudiantes analizados", total)
        k2.metric("Riesgo bajo", bajo)
        k3.metric("Riesgo moderado", mod)
        k4.metric("Riesgo alto", alto)
        cc1, cc2 = st.columns([1, 2])
        with cc1:
            fig = px.pie(dff, names="riesgo", title="Distribución por riesgo", color="riesgo", color_discrete_map=RIESGO_COLOR)
            st.plotly_chart(fig, use_container_width=True)
        with cc2:
            if "cambio" in dff.columns:
                fig2 = px.histogram(dff, x="cambio", color="riesgo", title="Evolución respecto al registro anterior", color_discrete_map=RIESGO_COLOR)
                st.plotly_chart(fig2, use_container_width=True)
        st.markdown("### Tabla de seguimiento")
        cols = [c for c in ["carne", "nombre", "correo", "curso", "seccion", "actividades_pct", "promedio", "entregas_tarde", "ingresos_semana", "dias_inactivo", "horas_respuesta", "riesgo", "riesgo_anterior", "cambio", "motivo_detectado", "asesor_bienestar"] if c in dff.columns]
        st.dataframe(dff[cols], use_container_width=True, height=420)

# -----------------------------------------------------------------------------
# Estudiante
# -----------------------------------------------------------------------------
with tabs[4]:
    df = st.session_state.analysis_df.copy()
    if df.empty:
        st.warning("No hay análisis activo.")
    else:
        names = (df["nombre"].fillna("") + " | " + df["carne"].fillna("").astype(str)).tolist()
        idx = st.selectbox("Seleccionar estudiante", range(len(names)), format_func=lambda i: names[i])
        row = df.iloc[idx]
        c1, c2, c3 = st.columns(3)
        c1.metric("Riesgo actual", row.get("riesgo", ""))
        c2.metric("Riesgo anterior", row.get("riesgo_anterior", "Sin registro"))
        c3.metric("Cambio", row.get("cambio", ""))
        st.write("**Motivo detectado:**", row.get("motivo_detectado", ""))
        st.dataframe(pd.DataFrame([row]), use_container_width=True)
        hist = st.session_state.db.get("Historial_Estudiantes", pd.DataFrame())
        if not hist.empty and row.get("carne") in hist.get("carne", pd.Series()).astype(str).values:
            st.markdown("### Historial registrado")
            st.dataframe(hist[hist["carne"].astype(str) == str(row.get("carne"))].tail(10), use_container_width=True)

# -----------------------------------------------------------------------------
# Mensajes
# -----------------------------------------------------------------------------
with tabs[5]:
    df = st.session_state.analysis_df.copy()
    if df.empty:
        st.warning("No hay análisis activo.")
    else:
        risk_msg = st.multiselect("Generar mensajes para riesgo", ["Bajo", "Moderado", "Alto"], default=["Moderado", "Alto"])
        candidates = df[df["riesgo"].isin(risk_msg)].copy()
        selected = st.multiselect("Seleccionar estudiantes", candidates.index.tolist(), format_func=lambda i: f"{candidates.loc[i,'nombre']} | {candidates.loc[i,'riesgo']}")
        if selected:
            for i in selected:
                row = candidates.loc[i]
                msg = generate_message(row, asesor_academico, horario_atencion, canal_atencion)
                with st.expander(f"{row.get('nombre')} - {row.get('riesgo')}"):
                    edited = st.text_area("Mensaje", msg, height=230, key=f"msg_{i}")
                    colx, coly = st.columns(2)
                    if colx.button("Registrar mensaje", key=f"regmsg_{i}"):
                        append_rows(st.session_state.db, "Mensajes_Enviados", [{"fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "carne": row.get("carne"), "nombre": row.get("nombre"), "correo": row.get("correo"), "curso": row.get("curso"), "riesgo": row.get("riesgo"), "tipo_mensaje": f"Seguimiento {row.get('riesgo')}", "mensaje_generado": edited, "enviado_canvas": "No", "asesor_academico": asesor_academico}])
                        st.success("Mensaje registrado en la base.")
                    if coly.button("Enviar por Canvas", key=f"send_{i}"):
                        client = st.session_state.canvas_client or (CanvasClient(url, token) if url and token else None)
                        if not client:
                            st.error("Primero valide Canvas o ingrese URL/token.")
                        else:
                            recip = str(row.get("canvas_user_id") or row.get("correo") or "").strip()
                            ok, resp = client.send_message([recip], f"Seguimiento académico - {row.get('curso')}", edited)
                            st.success(resp) if ok else st.error(resp)

# -----------------------------------------------------------------------------
# Derivaciones
# -----------------------------------------------------------------------------
with tabs[6]:
    df = st.session_state.analysis_df.copy()
    if df.empty:
        st.warning("No hay análisis activo.")
    else:
        derivables = df[df["riesgo"].isin(["Moderado", "Alto"])].copy()
        bienestar = st.session_state.db.get("Asesores_Bienestar", pd.DataFrame())
        advisor_names = bienestar["nombre"].dropna().tolist() if not bienestar.empty and "nombre" in bienestar.columns else []
        advisor = st.selectbox("Asesor de bienestar receptor", advisor_names + ["No especificado"])
        obs = st.text_area("Observaciones adicionales para los formatos", height=100)
        acciones = st.text_area("Acciones previas realizadas", value="Se revisó avance académico en Canvas, se identificó el nivel de riesgo, se generó mensaje de seguimiento y se registra el caso para acompañamiento oportuno.", height=100)
        selected = st.multiselect("Seleccionar estudiantes para derivar", derivables.index.tolist(), format_func=lambda i: f"{derivables.loc[i,'nombre']} | {derivables.loc[i,'riesgo']} | {derivables.loc[i].get('cambio','')}")
        if selected and st.button("Generar paquete de derivación", type="primary"):
            zip_buffer = io.BytesIO()
            report_rows = []
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                for i in selected:
                    row = derivables.loc[i].copy()
                    row["observaciones"] = obs
                    doc_bytes = make_derivation_doc(row, asesor_academico, advisor, obs, acciones)
                    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", str(row.get("nombre", "estudiante")))[:45]
                    zf.writestr(f"derivacion_{safe}_{row.get('riesgo')}.docx", doc_bytes)
                    prioridad = "Alta" if row.get("riesgo") == "Alto" else "Media"
                    report_rows.append({"id_derivacion": datetime.now().strftime("D%Y%m%d%H%M%S") + str(i), "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "carne": row.get("carne"), "nombre": row.get("nombre"), "correo": row.get("correo"), "curso": row.get("curso"), "seccion": row.get("seccion"), "riesgo": row.get("riesgo"), "prioridad": prioridad, "asesor_bienestar": advisor, "correo_bienestar": "", "motivo": row.get("motivo_detectado"), "acciones_previas": acciones, "observaciones": obs, "estado_derivacion": "Generada", "asesor_academico": asesor_academico})
                report_df = pd.DataFrame(report_rows)
                bio = io.BytesIO()
                with pd.ExcelWriter(bio, engine="xlsxwriter") as writer:
                    report_df.to_excel(writer, index=False, sheet_name="Listado_Derivaciones")
                zf.writestr("listado_general_derivaciones.xlsx", bio.getvalue())
            append_rows(st.session_state.db, "Derivaciones", report_rows)
            st.success("Paquete generado y derivaciones registradas.")
            st.download_button("Descargar paquete ZIP", zip_buffer.getvalue(), file_name=f"paquete_derivaciones_{date.today()}.zip", mime="application/zip")

# -----------------------------------------------------------------------------
# Export
# -----------------------------------------------------------------------------
with tabs[7]:
    st.markdown("### Exportar base actualizada")
    bytes_xlsx = export_db_excel(st.session_state.db)
    st.download_button("Descargar base de datos actualizada Excel", bytes_xlsx, file_name="base_datos_seguimiento_ave_actualizada.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if not st.session_state.analysis_df.empty:
        out = io.BytesIO()
        with pd.ExcelWriter(out, engine="xlsxwriter") as writer:
            st.session_state.analysis_df.to_excel(writer, index=False, sheet_name="Analisis_Actual")
        st.download_button("Descargar análisis actual", out.getvalue(), file_name="analisis_actual_estudiantes.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
